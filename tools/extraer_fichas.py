import re,json,collections
from pptx import Presentation
p=Presentation('/workspace/uploads/maraton/fichas.pptx')
out=[];hdr=None
for s in p.slides:
  for sh in s.shapes:
    if not sh.has_text_frame: continue
    t=sh.text_frame.text.strip()
    m=re.match(r'(.+?)\s{2,}([A-Z]{3})\s*(\d+)/(\d+)$',t)
    if m: hdr=(m.group(1).strip(),m.group(2),m.group(3));continue
    if hdr and re.match(r'1\.',t):
      for q in re.split(r'\n(?=\d+\.\s)',t):
        mm=re.match(r'\d+\.\s*(.*?)\n(.*)',q,re.S)
        if not mm: continue
        opts=re.findall(r'([ABC])\)\s*(.*?)(?=\s+[ABC]\)|\n|$)',mm.group(2))
        if len(opts)!=3: print('SKIP',hdr,q[:80]);continue
        ok=[i for i,(_,o) in enumerate(opts) if o.strip().endswith('*')]
        if len(ok)!=1: print('SKIP*',hdr,q[:80]);continue
        out.append({'tema':hdr[1],'nombreTema':hdr[0],'ficha':int(hdr[2]),'p':mm.group(1).strip(),
          'op':[o.strip().rstrip('*').strip() for _,o in opts],'ok':ok[0]})
      hdr=None
print(len(out),collections.Counter(q['tema'] for q in out),collections.Counter(q['nombreTema'] for q in out))
json.dump(out,open('data/preguntas.json','w'),ensure_ascii=False,indent=0)
